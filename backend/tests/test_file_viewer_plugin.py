from fastapi.testclient import TestClient
from main import app
from services.file_service import FileService


def test_file_content_and_raw_endpoints(tmp_path, monkeypatch):
    test_files_root = tmp_path / "files_root"
    test_files_root.mkdir()

    # Création de fichiers de test
    text_file = test_files_root / "hello.txt"
    text_file.write_text("Hello World!\nCeci est un test.", encoding="utf-8")

    code_file = test_files_root / "script.py"
    code_file.write_text("print('hello')", encoding="utf-8")

    json_file = test_files_root / "data.json"
    json_file.write_text('{"status": "ok", "count": 42}', encoding="utf-8")

    bin_file = test_files_root / "image.png"
    bin_file.write_bytes(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR")

    # Patcher la racine FileService
    monkeypatch.setattr(FileService, "root", test_files_root)

    with TestClient(app) as client:
        # 1. Lister les fichiers
        res = client.get("/api/files/list")
        assert res.status_code == 200
        filenames = [f["name"] for f in res.json()["files"]]
        assert "hello.txt" in filenames
        assert "script.py" in filenames
        assert "image.png" in filenames

        # 2. Contenu texte
        res = client.get("/api/files/content?path=hello.txt")
        assert res.status_code == 200
        data = res.json()
        assert data["name"] == "hello.txt"
        assert data["type"] == "text"
        assert "Hello World!" in data["content"]

        # 3. Contenu code python
        res = client.get("/api/files/content?path=script.py")
        assert res.status_code == 200
        data = res.json()
        assert data["type"] == "text"
        assert data["content"] == "print('hello')"

        # 4. Fichier image (métadonnées)
        res = client.get("/api/files/content?path=image.png")
        assert res.status_code == 200
        data = res.json()
        assert data["type"] == "image"

        # 5. Fichier raw / stream
        res = client.get("/api/files/raw?path=image.png")
        assert res.status_code == 200
        assert res.content.startswith(b"\x89PNG")

        # 6. Fichier introuvable
        res = client.get("/api/files/content?path=nonexistent.txt")
        assert res.status_code == 404

        # 7. Téléchargement avec header attachment
        res = client.get("/api/files/raw?path=hello.txt&download=true")
        assert res.status_code == 200
        assert 'attachment; filename="hello.txt"' in res.headers.get("content-disposition", "")


def test_office_documents_preview(tmp_path, monkeypatch):
    """Aperçu Word/Excel/PowerPoint : vérifie l'extraction structurée du contenu."""
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    test_files_root = tmp_path / "files_root"
    test_files_root.mkdir()

    doc = Document()
    doc.add_heading("Titre", level=1)
    doc.add_paragraph("Un paragraphe.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    doc.save(test_files_root / "doc.docx")

    wb = Workbook()
    ws = wb.active
    ws.append(["Nom", "Age"])
    ws.append(["Alice", 30])
    wb.save(test_files_root / "sheet.xlsx")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Slide 1"
    slide.placeholders[1].text = "Contenu"
    prs.save(test_files_root / "slides.pptx")

    monkeypatch.setattr(FileService, "root", test_files_root)

    with TestClient(app) as client:
        res = client.get("/api/files/content?path=doc.docx")
        assert res.status_code == 200
        data = res.json()
        assert data["type"] == "docx"
        assert data["document"]["paragraphs"][0]["text"] == "Titre"
        assert data["document"]["paragraphs"][0]["heading"] == 1
        assert data["document"]["tables"][0]["rows"] == [["A", "B"]]

        res = client.get("/api/files/content?path=sheet.xlsx")
        assert res.status_code == 200
        data = res.json()
        assert data["type"] == "xlsx"
        assert data["workbook"]["sheets"][0]["rows"][0] == ["Nom", "Age"]
        assert data["workbook"]["sheets"][0]["rows"][1] == ["Alice", "30"]

        res = client.get("/api/files/content?path=slides.pptx")
        assert res.status_code == 200
        data = res.json()
        assert data["type"] == "pptx"
        assert data["presentation"]["slides"][0]["title"] == "Slide 1"
        assert data["presentation"]["slides"][0]["texts"] == ["Contenu"]


def test_sqlite_database_preview(tmp_path, monkeypatch):
    """Aperçu d'une base SQLite : vérifie la liste des tables, colonnes et lignes."""
    import sqlite3

    test_files_root = tmp_path / "files_root"
    test_files_root.mkdir()

    db_path = test_files_root / "app.db"
    conn = sqlite3.connect(db_path)
    conn.execute("CREATE TABLE users (id INTEGER PRIMARY KEY, name TEXT, age INTEGER)")
    conn.execute("INSERT INTO users (name, age) VALUES ('Alice', 30)")
    conn.execute("INSERT INTO users (name, age) VALUES ('Bob', 25)")
    conn.commit()
    conn.close()

    monkeypatch.setattr(FileService, "root", test_files_root)

    with TestClient(app) as client:
        res = client.get("/api/files/content?path=app.db")
        assert res.status_code == 200
        data = res.json()
        assert data["type"] == "db"
        tables = data["database"]["tables"]
        assert len(tables) == 1
        assert tables[0]["name"] == "users"
        assert tables[0]["columns"] == ["id", "name", "age"]
        assert tables[0]["rows"] == [["1", "Alice", "30"], ["2", "Bob", "25"]]
        assert tables[0]["row_count"] == 2

        # Fichier non-SQLite avec extension .db : doit retomber en binaire (téléchargement).
        fake_db = test_files_root / "fake.db"
        fake_db.write_bytes(b"not a sqlite file")
        res = client.get("/api/files/content?path=fake.db")
        assert res.status_code == 200
        assert res.json()["type"] == "binary"


def test_search_folder_file_rename_delete_endpoints(tmp_path, monkeypatch):
    """Nouveaux endpoints pilotables par l'assistant fichiers (chat/vocal) :
    recherche par motif glob, création de dossier/fichier, renommage, suppression."""
    test_files_root = tmp_path / "files_root"
    test_files_root.mkdir()
    (test_files_root / "notes.txt").write_text("un fichier existant", encoding="utf-8")
    (test_files_root / "report.pdf").write_bytes(b"%PDF-1.4")

    monkeypatch.setattr(FileService, "root", test_files_root)

    with TestClient(app) as client:
        # Recherche glob
        res = client.get("/api/files/search?pattern=*.txt")
        assert res.status_code == 200
        matches = res.json()["matches"]
        assert "notes.txt" in matches
        assert "report.pdf" not in matches

        # Création de dossier
        res = client.post("/api/files/folder", json={"path": "docs"})
        assert res.status_code == 200
        assert (test_files_root / "docs").is_dir()

        # Création de dossier déjà existant -> conflit
        res = client.post("/api/files/folder", json={"path": "docs"})
        assert res.status_code == 409

        # Création de fichier avec contenu
        res = client.post("/api/files/file", json={"path": "docs/todo.txt", "content": "acheter du pain"})
        assert res.status_code == 200
        assert (test_files_root / "docs" / "todo.txt").read_text(encoding="utf-8") == "acheter du pain"

        # Renommage
        res = client.post("/api/files/rename", json={"path": "docs/todo.txt", "new_name": "liste.txt"})
        assert res.status_code == 200
        assert (test_files_root / "docs" / "liste.txt").exists()
        assert not (test_files_root / "docs" / "todo.txt").exists()

        # Renommage d'un fichier inexistant -> 404
        res = client.post("/api/files/rename", json={"path": "docs/absent.txt", "new_name": "x.txt"})
        assert res.status_code == 404

        # Suppression
        res = client.delete("/api/files/delete?path=docs/liste.txt")
        assert res.status_code == 200
        assert not (test_files_root / "docs" / "liste.txt").exists()

        # Suppression de la racine interdite
        res = client.delete("/api/files/delete?path=.")
        assert res.status_code == 400


def test_summary_endpoint(tmp_path, monkeypatch):
    """Compte rendu récursif d'un dossier (nombre de fichiers, taille, répartition par
    extension) — utilisé par l'assistant fichiers pour répondre à "fais-moi un compte rendu"."""
    test_files_root = tmp_path / "files_root"
    test_files_root.mkdir()
    (test_files_root / "a.txt").write_text("hello")
    (test_files_root / "sub").mkdir()
    (test_files_root / "sub" / "b.py").write_text("print(1)")

    monkeypatch.setattr(FileService, "root", test_files_root)

    with TestClient(app) as client:
        res = client.get("/api/files/summary")
        assert res.status_code == 200
        data = res.json()
        assert data["file_count"] == 2
        assert data["folder_count"] == 1
        assert any(item["extension"] == ".py" for item in data["by_extension"])

        # Chemin qui n'est pas un dossier -> 400
        res = client.get("/api/files/summary?path=a.txt")
        assert res.status_code == 400


def test_file_chat_handler_matches():
    """Détection des demandes en langage naturel destinées à l'assistant fichiers du chat/vocal
    (backend/plugins/files/chat_handler.py) — pas d'appel réel à Claude ici."""
    from plugins.files.chat_handler import matches

    assert matches("ouvre le fichier rapport.pdf")
    assert matches("trouve-moi le fichier budget dans mes documents")
    assert matches("renomme le dossier Photos en Vacances")
    assert matches("supprime ce fichier")
    assert matches("filtre les fichiers par *.csv")
    assert not matches("quel temps fait-il aujourd'hui ?")
    assert not matches("raconte-moi une blague")


def test_file_assistant_accueil_alias(tmp_path, monkeypatch):
    """« Accueil » est le raccourci affiché dans l'interface pour la racine (path=".") : l'assistant
    fichiers doit le reconnaître au lieu de chercher un dossier littéralement nommé "accueil"."""
    import asyncio

    from services import file_assistant

    test_files_root = tmp_path / "files_root"
    test_files_root.mkdir()
    (test_files_root / "a.txt").write_text("hello")
    monkeypatch.setattr(FileService, "root", test_files_root)

    async def run_all():
        nav = {}
        result = await file_assistant._execute_tool("list_directory", {"path": "Accueil"}, nav)
        assert nav["path"] == "."
        assert result["count"] == 1

        nav2 = {}
        summary = await file_assistant._execute_tool("summarize_directory", {"path": "ACCUEIL"}, nav2)
        assert nav2["path"] == "."
        assert summary["file_count"] == 1

    asyncio.run(run_all())
